# app.py — Notes Compiler (OpenAI API) + Scanned PDF OCR
# ✅ Compile → Download (NO Q&A / chat)
#
# Install:
#   pip install streamlit openai pypdf python-docx python-pptx pillow pymupdf
#
# Run (PowerShell):
#   $env:OPENAI_API_KEY="sk-..."
#   # optional upload limit (MB):
#   $env:STREAMLIT_SERVER_MAX_UPLOAD_SIZE="2000"
#   streamlit run app.py

import os
import io
import re
import base64
import hashlib
from dataclasses import dataclass
from typing import List, Optional, Tuple

import streamlit as st

# Optional deps
try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    import docx  # python-docx
except Exception:
    docx = None

try:
    from pptx import Presentation  # python-pptx
except Exception:
    Presentation = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None

try:
    from openai import OpenAI
except Exception:
    OpenAI = None


# -----------------------------
# UI CONFIG
# -----------------------------
st.set_page_config(page_title="Notes Compiler", page_icon="🧠", layout="wide")

CSS = """
<style>
.block-container { padding-top: 1.2rem; padding-bottom: 3rem; }
h1, h2, h3 { letter-spacing: -0.02em; }
.stButton>button, .stDownloadButton>button { border-radius: 14px; padding: 0.6rem 1rem; }
code { border-radius: 10px; padding: 0.15rem 0.35rem; }
small { opacity: 0.75; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

st.title("🧠 Notes Compiler (Compile → Download)")
st.caption("Upload notes (PDF/DOCX/PPTX/TXT/MD + images). Scanned PDFs are OCR’d via OpenAI vision.")


# -----------------------------
# DATA MODEL
# -----------------------------
@dataclass
class DocItem:
    name: str
    kind: str         # 'text' | 'pdf' | 'docx' | 'pptx' | 'image' | 'unknown'
    mime: str
    sha256: str
    text: Optional[str] = None
    image_bytes: Optional[bytes] = None
    raw_bytes: Optional[bytes] = None  # keep PDF bytes for OCR


# -----------------------------
# HELPERS
# -----------------------------
def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def clean_text(s: str) -> str:
    if not s:
        return ""
    s = s.replace("\x00", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def approx_tokens(text: str) -> int:
    # rough: ~4 chars/token
    if not text:
        return 0
    return max(1, int(len(text) / 4))


def split_paragraphs(text: str) -> List[str]:
    parts = re.split(r"\n\s*\n", (text or "").strip())
    return [p.strip() for p in parts if p.strip()]


def chunk_text(text: str, target_tokens: int = 1800, overlap_tokens: int = 200) -> List[str]:
    text = clean_text(text)
    paras = split_paragraphs(text)
    if not paras:
        return []

    chunks: List[str] = []
    cur: List[str] = []
    cur_tokens = 0

    def flush():
        nonlocal cur, cur_tokens
        if cur:
            chunk = "\n\n".join(cur).strip()
            if chunk:
                chunks.append(chunk)
        cur = []
        cur_tokens = 0

    for p in paras:
        pt = approx_tokens(p)

        # Very large paragraph → split by sentences
        if pt > target_tokens:
            flush()
            sentences = re.split(r"(?<=[.!?])\s+", p.strip())
            buf: List[str] = []
            bt = 0
            for sent in sentences:
                stoks = approx_tokens(sent)
                if bt + stoks > target_tokens and buf:
                    chunks.append(" ".join(buf).strip())
                    if overlap_tokens > 0:
                        tail = " ".join(buf)[-overlap_tokens * 4 :].strip()
                        buf = [tail] if tail else []
                        bt = approx_tokens(" ".join(buf))
                    else:
                        buf = []
                        bt = 0
                buf.append(sent)
                bt += stoks
            if buf:
                chunks.append(" ".join(buf).strip())
            continue

        # Normal case
        if cur_tokens + pt > target_tokens and cur:
            flush()
            if overlap_tokens > 0 and chunks:
                tail = chunks[-1][-overlap_tokens * 4 :].strip()
                if tail:
                    cur = [tail]
                    cur_tokens = approx_tokens(tail)

        cur.append(p)
        cur_tokens += pt

    flush()
    return chunks


def get_openai_client() -> "OpenAI":
    if OpenAI is None:
        raise RuntimeError("OpenAI SDK not installed. Run: pip install openai")
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("Missing OPENAI_API_KEY env var (set it in your terminal).")
    return OpenAI(api_key=api_key)


def response_to_text(resp) -> str:
    t = getattr(resp, "output_text", None)
    if isinstance(t, str) and t.strip():
        return t.strip()

    out = []
    items = getattr(resp, "output", []) or []
    for item in items:
        if getattr(item, "type", None) == "message":
            content = getattr(item, "content", []) or []
            for c in content:
                ctype = getattr(c, "type", None)
                if ctype in ("output_text", "text"):
                    txt = getattr(c, "text", None)
                    if isinstance(txt, str) and txt.strip():
                        out.append(txt.strip())
    return "\n\n".join(out).strip()


def oai_text(
    prompt: str,
    *,
    model: str,
    instructions: str,
    temperature: float,
    max_output_tokens: int,
    store: bool,
) -> str:
    client = get_openai_client()
    resp = client.responses.create(
        model=model,
        instructions=instructions,
        input=prompt,
        temperature=temperature,
        max_output_tokens=max_output_tokens,
        store=store,
    )
    return response_to_text(resp)


def oai_vision(
    prompt: str,
    image_bytes: bytes,
    mime: str,
    *,
    model: str,
    instructions: str,
    temperature: float,
    max_output_tokens: int,
    store: bool,
) -> str:
    client = get_openai_client()
    b64 = base64.b64encode(image_bytes).decode("utf-8")
    data_url = f"data:{mime};base64,{b64}"

    input_items = [
        {
            "role": "user",
            "content": [
                {"type": "input_text", "text": prompt},
                {"type": "input_image", "image_url": data_url},
            ],
        }
    ]

    resp = client.responses.create(
        model=model,
        instructions=instructions,
        input=input_items,
        temperature=temperature,
        max_output_tokens=max_output_tokens,
        store=store,
    )
    return response_to_text(resp)


# -----------------------------
# FILE EXTRACTION
# -----------------------------
def extract_text_from_upload(name: str, mime: str, data: bytes) -> Tuple[str, str]:
    ext = (name.split(".")[-1] or "").lower()

    if ext in ("txt", "md", "markdown", "csv", "log"):
        try:
            return clean_text(data.decode("utf-8")), "text"
        except Exception:
            return clean_text(data.decode("latin-1", errors="ignore")), "text"

    if ext == "pdf":
        if PdfReader is None:
            return "", "pdf"
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = []
            for p in reader.pages:
                txt = p.extract_text() or ""
                if txt.strip():
                    pages.append(txt)
            return clean_text("\n\n".join(pages)), "pdf"
        except Exception:
            return "", "pdf"

    if ext == "docx":
        if docx is None:
            return "", "docx"
        try:
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
            return clean_text("\n".join(parts)), "docx"
        except Exception:
            return "", "docx"

    if ext == "pptx":
        if Presentation is None:
            return "", "pptx"
        try:
            prs = Presentation(io.BytesIO(data))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        t = shape.text.strip()
                        if t:
                            parts.append(t)
            return clean_text("\n".join(parts)), "pptx"
        except Exception:
            return "", "pptx"

    if ext in ("png", "jpg", "jpeg", "webp", "bmp"):
        return "", "image"

    return "", "unknown"


# -----------------------------
# OCR HELPERS (Scanned PDFs)
# -----------------------------
def is_text_sparse(text: Optional[str], threshold_chars: int) -> bool:
    if not text:
        return True
    return len(text.strip()) < threshold_chars


def maybe_downscale_png(png_bytes: bytes, max_dim: int) -> bytes:
    if Image is None:
        return png_bytes
    try:
        img = Image.open(io.BytesIO(png_bytes))
        w, h = img.size
        if max(w, h) <= max_dim:
            return png_bytes
        scale = max_dim / float(max(w, h))
        new_size = (max(1, int(w * scale)), max(1, int(h * scale)))
        img = img.resize(new_size)
        out = io.BytesIO()
        img.save(out, format="PNG", optimize=True)
        return out.getvalue()
    except Exception:
        return png_bytes


def pdf_pages_to_pngs(pdf_bytes: bytes, zoom: float, max_pages: int):
    if fitz is None:
        raise RuntimeError("PyMuPDF not installed. Run: pip install pymupdf")
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    n = doc.page_count
    if max_pages and max_pages > 0:
        n = min(n, max_pages)

    mat = fitz.Matrix(zoom, zoom)
    for i in range(n):
        page = doc.load_page(i)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        yield (i + 1), pix.tobytes("png")


# -----------------------------
# PROMPTS
# -----------------------------
BASE_INSTRUCTIONS = """You are a strict, helpful NOTES COMPILER.
Rules:
- Be accurate. If unclear, label it as [CHECK] instead of guessing.
- Prefer structured bullet points and short paragraphs.
- Keep formulas, definitions, and step-by-step processes.
- Remove fluff, keep exam-relevant detail.
- Use Markdown headings.
"""

def make_chunk_prompt(doc_name: str, chunk_idx: int, chunk_text: str, mode: str, level: str) -> str:
    return f"""
You are summarizing chunk {chunk_idx} from document: {doc_name}

OUTPUT STYLE: {level}
MODE: {mode}

Chunk text:
\"\"\"\n{chunk_text}\n\"\"\"

Write tight study notes in Markdown with:
- Key ideas (bullets)
- Definitions (if any)
- Processes/steps (if any)
- Formulas/equations (if any)
- Worked micro-example (if the chunk supports it)
- Common mistakes / pitfalls
- 3 quick recall questions + answers (short)

If something is missing in the chunk, do not invent it.
"""


def make_doc_merge_prompt(doc_name: str, chunk_summaries: List[str], mode: str, level: str) -> str:
    joined = "\n\n---\n\n".join(chunk_summaries)
    return f"""
You will merge chunk summaries into ONE clean document note.

Document: {doc_name}
OUTPUT STYLE: {level}
MODE: {mode}

Chunk summaries:
{joined}

Produce a single cohesive Markdown note with:
- A mini table of contents
- Clear headings
- No duplicates
- Consistent terminology
- A short "Quick Revision" section at the top (10 bullets max)
"""


def make_global_merge_prompt(all_doc_notes: List[Tuple[str, str]], mode: str, level: str) -> str:
    body = "\n\n".join([f"## Source: {name}\n\n{note}" for name, note in all_doc_notes])
    return f"""
You will compile MULTIPLE document notes into ONE master set of notes.

OUTPUT STYLE: {level}
MODE: {mode}

Requirements:
- Start with: Title, then a Table of Contents
- Then sections grouped by topic (merge overlaps even if different files)
- Keep it exam-ready: definitions, steps, formulas, examples, pitfalls
- Add a final "1-page Cheat Sheet" section (dense bullets)

Here are the source notes:
{body}
"""


# -----------------------------
# SIDEBAR SETTINGS
# -----------------------------
with st.sidebar:
    st.header("⚙️ Settings")

    if os.getenv("OPENAI_API_KEY", "").strip():
        st.success("OPENAI_API_KEY detected ✅")
    else:
        st.warning("No OPENAI_API_KEY found. Set it in your terminal.")

    model = st.text_input("Model", value="gpt-4o-mini")
    store = st.toggle("Store responses on OpenAI (privacy)", value=False)
    temperature = st.slider("Creativity (temperature)", 0.0, 1.2, 0.2, 0.05)

    st.divider()
    mode = st.selectbox("Compile mode", ["Study Notes", "Exam Revision", "Lecture Cleanup", "Flashcards-first"])
    level = st.selectbox("Depth", ["Simple (high school)", "Standard (A-level)", "Advanced (university)"], index=1)

    st.divider()
    chunk_tokens = st.slider("Chunk size (approx tokens)", 800, 4000, 1800, 100)
    overlap_tokens = st.slider("Chunk overlap (approx tokens)", 0, 600, 200, 25)

    st.divider()
    st.subheader("🧾 Scanned PDF OCR")
    enable_pdf_ocr = st.toggle("OCR scanned PDFs", value=True)
    pdf_ocr_chars_threshold = st.number_input("Treat PDF as scanned if extracted text < (chars)", 0, 10000, 300, 50)
    pdf_ocr_max_pages = st.number_input("Max PDF pages to OCR (0 = all)", 0, 10000, 30, 5)
    pdf_render_zoom = st.slider("PDF render quality (zoom)", 1.0, 3.0, 2.0, 0.25)
    ocr_downscale_max_dim = st.number_input("Downscale OCR images max dimension (px)", 800, 5000, 2200, 100)

    if enable_pdf_ocr and fitz is None:
        st.warning("OCR needs PyMuPDF: pip install pymupdf")


# -----------------------------
# SESSION STATE
# -----------------------------
for key, default in [
    ("docs", {}),
    ("chunk_cache", {}),
    ("doc_cache", {}),
    ("compiled_notes", None),
    ("doc_notes_list", []),
    ("compiled_key", None),
]:
    if key not in st.session_state:
        st.session_state[key] = default


# -----------------------------
# MAIN
# -----------------------------
uploaded = st.file_uploader(
    "Upload notes (PDF/DOCX/PPTX/TXT/MD + images). Multiple files supported.",
    type=["pdf", "docx", "pptx", "txt", "md", "markdown", "png", "jpg", "jpeg", "webp", "bmp"],
    accept_multiple_files=True,
)

c1, c2 = st.columns([1, 1], vertical_alignment="top")
with c1:
    compile_btn = st.button("🚀 Compile", type="primary", use_container_width=True)
with c2:
    clear_btn = st.button("🧹 Clear", use_container_width=True)

if clear_btn:
    for k in ["docs", "chunk_cache", "doc_cache", "compiled_notes", "doc_notes_list", "compiled_key"]:
        st.session_state.pop(k, None)
    st.toast("Cleared.", icon="✅")
    st.stop()


def ingest_files(files) -> List[DocItem]:
    items: List[DocItem] = []
    for f in files or []:
        data = f.getvalue()
        sha = sha256_bytes(data)
        name = f.name
        mime = getattr(f, "type", "") or "application/octet-stream"

        if sha in st.session_state["docs"]:
            items.append(st.session_state["docs"][sha])
            continue

        text, kind = extract_text_from_upload(name, mime, data)
        doc_item = DocItem(
            name=name,
            kind=kind,
            mime=mime,
            sha256=sha,
            text=text if text else None,
            image_bytes=data if kind == "image" else None,
            raw_bytes=data if kind == "pdf" else None,
        )
        st.session_state["docs"][sha] = doc_item
        items.append(doc_item)
    return items


docs = ingest_files(uploaded)

if docs:
    total_chars = sum(len(d.text or "") for d in docs)
    st.info(f"Loaded {len(docs)} file(s). Extracted text chars so far: {total_chars:,}")

    low_text_pdfs = [d for d in docs if d.kind == "pdf" and is_text_sparse(d.text, int(pdf_ocr_chars_threshold))]
    if low_text_pdfs:
        st.warning("Some PDFs look scanned / low-text. OCR will run (if enabled) and then they’ll compile normally.")


def compile_pipeline(docs: List[DocItem]) -> str:
    if not docs:
        raise RuntimeError("No files uploaded.")

    # 1) Transcribe uploaded images into text (and lightly structure them)
    vision_docs = [d for d in docs if d.kind == "image"]
    if vision_docs:
        st.subheader("🖼️ Transcribing uploaded images")
        for i, d in enumerate(vision_docs, 1):
            st.write(f"Image {i}/{len(vision_docs)}: `{d.name}`")
            cache_key = f"img:{d.sha256}:{model}:{level}:{mode}"
            if cache_key in st.session_state["doc_cache"]:
                d.text = st.session_state["doc_cache"][cache_key]
                st.caption("Cached ✅")
                continue

            prompt = "Extract ALL readable text from this image of notes. Output as plain text (no summarizing)."
            txt = oai_vision(
                prompt,
                d.image_bytes or b"",
                d.mime or "image/png",
                model=model,
                instructions="You are an OCR engine. Output text only.",
                temperature=0.0,
                max_output_tokens=2000,
                store=store,
            )
            d.text = clean_text(txt)
            st.session_state["doc_cache"][cache_key] = d.text

    # 2) OCR scanned PDFs
    if enable_pdf_ocr:
        scanned = [d for d in docs if d.kind == "pdf" and is_text_sparse(d.text, int(pdf_ocr_chars_threshold))]
        if scanned:
            st.subheader("🧾 OCR for scanned / low-text PDFs")
            if fitz is None:
                raise RuntimeError("OCR enabled but PyMuPDF is missing. Run: pip install pymupdf")

            for d in scanned:
                if not d.raw_bytes:
                    continue

                st.write(f"OCR: `{d.name}`")
                page_texts: List[str] = []

                for page_no, png_bytes in pdf_pages_to_pngs(
                    d.raw_bytes,
                    zoom=float(pdf_render_zoom),
                    max_pages=int(pdf_ocr_max_pages),
                ):
                    png_bytes = maybe_downscale_png(png_bytes, max_dim=int(ocr_downscale_max_dim))
                    cache_key = f"pdfocr:{d.sha256}:p{page_no}:z{pdf_render_zoom}:md{ocr_downscale_max_dim}:{model}"

                    if cache_key in st.session_state["chunk_cache"]:
                        page_texts.append(st.session_state["chunk_cache"][cache_key])
                        continue

                    prompt = (
                        "Extract ALL readable text from this scanned PDF page. "
                        "Do NOT summarize. Preserve headings, bullets, numbering, and equations."
                    )
                    page_txt = oai_vision(
                        prompt,
                        png_bytes,
                        "image/png",
                        model=model,
                        instructions="You are an OCR engine. Output text only. No extra commentary.",
                        temperature=0.0,
                        max_output_tokens=2000,
                        store=store,
                    )
                    page_txt = clean_text(page_txt)
                    st.session_state["chunk_cache"][cache_key] = page_txt
                    page_texts.append(page_txt)

                d.text = clean_text("\n\n".join(page_texts))

    # 3) Chunk + summarize each doc
    st.subheader("🧩 Chunking & summarizing")
    usable_docs = [d for d in docs if (d.text or "").strip()]
    if not usable_docs:
        raise RuntimeError("No usable text extracted. (If scanned PDFs: enable OCR and install pymupdf.)")

    progress = st.progress(0.0)
    status = st.empty()

    # Estimate progress steps
    doc_chunks_map = {}
    total_steps = 0
    for d in usable_docs:
        chs = chunk_text(d.text or "", int(chunk_tokens), int(overlap_tokens))
        doc_chunks_map[d.sha256] = chs
        total_steps += max(1, len(chs))
    total_steps = max(1, total_steps)

    done = 0
    doc_notes_list: List[Tuple[str, str]] = []

    for d in usable_docs:
        chunks = doc_chunks_map.get(d.sha256, [])
        if not chunks:
            continue

        chunk_summaries: List[str] = []
        for idx, ch in enumerate(chunks, 1):
            done += 1
            status.write(f"Summarizing `{d.name}` chunk {idx}/{len(chunks)} …")
            progress.progress(min(1.0, done / total_steps))

            ck = f"chunk:{d.sha256}:{idx}:{model}:{level}:{mode}:{chunk_tokens}:{overlap_tokens}"
            if ck in st.session_state["chunk_cache"]:
                chunk_summaries.append(st.session_state["chunk_cache"][ck])
                continue

            prompt = make_chunk_prompt(d.name, idx, ch, mode, level)
            summ = oai_text(
                prompt,
                model=model,
                instructions=BASE_INSTRUCTIONS,
                temperature=temperature,
                max_output_tokens=1500,
                store=store,
            )
            summ = clean_text(summ)
            st.session_state["chunk_cache"][ck] = summ
            chunk_summaries.append(summ)

        # Merge per document
        status.write(f"Merging chunks for `{d.name}` …")
        mk = f"docmerge:{d.sha256}:{model}:{level}:{mode}"
        if mk in st.session_state["doc_cache"]:
            merged = st.session_state["doc_cache"][mk]
        else:
            prompt = make_doc_merge_prompt(d.name, chunk_summaries, mode, level)
            merged = oai_text(
                prompt,
                model=model,
                instructions=BASE_INSTRUCTIONS,
                temperature=max(0.0, min(0.35, temperature)),
                max_output_tokens=2200,
                store=store,
            )
            merged = clean_text(merged)
            st.session_state["doc_cache"][mk] = merged

        doc_notes_list.append((d.name, merged))

    if not doc_notes_list:
        raise RuntimeError("Nothing to compile after processing.")

    st.session_state["doc_notes_list"] = doc_notes_list

    # 4) Global compile
    st.subheader("🧱 Building compiled notes")
    compiled_key = f"{model}|{level}|{mode}|" + "|".join(sorted([d.sha256 for d in usable_docs]))
    if st.session_state.get("compiled_notes") and st.session_state.get("compiled_key") == compiled_key:
        return st.session_state["compiled_notes"]

    prompt = make_global_merge_prompt(doc_notes_list, mode, level)
    compiled = oai_text(
        prompt,
        model=model,
        instructions=BASE_INSTRUCTIONS,
        temperature=max(0.0, min(0.35, temperature)),
        max_output_tokens=3200,
        store=store,
    )
    compiled = clean_text(compiled)

    st.session_state["compiled_key"] = compiled_key
    st.session_state["compiled_notes"] = compiled
    return compiled


# -----------------------------
# RUN COMPILE
# -----------------------------
if compile_btn:
    try:
        if OpenAI is None:
            st.error("Missing OpenAI SDK. Install: pip install openai")
        else:
            compiled = compile_pipeline(docs)
            st.success("Compiled ✅ Scroll down to download.")
    except Exception as e:
        st.error(f"Compile failed: {e}")


# -----------------------------
# OUTPUT + DOWNLOAD
# -----------------------------
compiled = st.session_state.get("compiled_notes")
if compiled:
    st.subheader("📘 Compiled Notes (Preview)")
    st.markdown(compiled)

    data = compiled.encode("utf-8")
    a, b = st.columns([1, 1], vertical_alignment="center")
    with a:
        st.download_button("⬇️ Download Markdown", data=data, file_name="compiled_notes.md", mime="text/markdown")
    with b:
        st.download_button("⬇️ Download TXT", data=data, file_name="compiled_notes.txt", mime="text/plain")

    with st.expander("🔎 Per-file merged notes (optional)", expanded=False):
        for name, note in st.session_state.get("doc_notes_list", []):
            st.markdown(f"### {name}")
            st.markdown(note)
            st.divider()
